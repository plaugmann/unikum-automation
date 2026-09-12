"""Udtraek ren tekst fra beskeder og bilag.

Laererne blander formater: noget staar i selve beskeden, noget i en vedhaeftet
PDF, Word-fil eller txt. LLM'en skal have det hele som ren tekst, uanset hvor
det kom fra.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

from bs4 import BeautifulSoup

MAX_ATTACHMENT_CHARS = 40_000

# Laererne laegger ogsaa skaermbilleder op. Dem kan vi ikke traekke tekst ud
# af, men LLM'en kan se dem, saa de sendes videre som billeder.
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp"}


def is_image(filename: str) -> bool:
    return Path(filename).suffix.lower() in IMAGE_SUFFIXES


def html_to_text(html: str | None) -> str:
    """Unikums brodtekst er HTML fra en draft-editor."""
    if not html:
        return ""
    soup = BeautifulSoup(html, "lxml")

    # Bilags-elementer staar som links i teksten. Selve filen henter vi
    # separat, saa her erstatter vi dem med en laesbar markoer.
    for a in soup.select('[data-custom-entity-type="INTERNAL_ATTACHMENT"]'):
        name = a.get("data-file-name") or a.get_text(strip=True) or "bilag"
        a.replace_with(f"[Bilag: {name}]")

    for br in soup.find_all("br"):
        br.replace_with("\n")
    for block in soup.find_all(["p", "div", "li", "tr", "h1", "h2", "h3", "h4"]):
        block.append("\n")

    text = soup.get_text()
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def pdf_to_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = [(page.extract_text() or "") for page in reader.pages]
    return "\n\n".join(p.strip() for p in parts if p.strip()).strip()


def docx_to_text(data: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(p for p in parts if p.strip()).strip()


def plain_to_text(data: bytes) -> str:
    for encoding in ("utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace").strip()


def attachment_to_text(filename: str, data: bytes) -> str:
    """Vaelg udtraekker ud fra filendelsen. Kaster ved ukendt eller defekt fil."""
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text = pdf_to_text(data)
    elif suffix in (".docx", ".docm"):
        text = docx_to_text(data)
    elif suffix in (".txt", ".csv", ".md"):
        text = plain_to_text(data)
    elif suffix == ".doc":
        # Gammelt binaert Word-format. Sjaeldent nok til at vi hellere vil
        # vide det end at traekke en tung afhaengighed ind.
        raise ValueError("gammelt .doc-format understoettes ikke")
    else:
        raise ValueError(f"ukendt filtype: {suffix or 'ingen endelse'}")

    if not text:
        raise ValueError("ingen tekst kunne udtraekkes (scannet dokument?)")
    return text[:MAX_ATTACHMENT_CHARS]


def filename_from_path(path: str) -> str:
    """Unikums bilagsstier ser ud som /content/4cc0/<id>-<uuid>.pdf"""
    return Path(path).name
